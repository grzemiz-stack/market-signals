"""Etap 4N: etykiety przyczyn — blokada PIT, parsery źródeł, duże ruchy, dopasowanie kalendarza, arkusz E2."""

import io
import json
import re
import sys
from datetime import date
from pathlib import Path

import duckdb
import numpy as np
import pandas as pd
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from ms import causes, db  # noqa: E402
from ms import cause_sources as cs  # noqa: E402

CAUSE_TABLES = re.compile(r"\b(move_causes|cause_events|bbds_codings)\b")
ALLOWED = {"ms/db.py", "ms/causes.py", "ms/cause_sources.py", "build_causes.py",
           "ms/views.py",   # 4C: wyświetlanie „tego dnia było” (describe_day) — nie cecha; importerzy pilnowani niżej
           "ms/classify.py", "validate_causes.py"}   # 4N: klasyfikator i jego walidacja (BBDS = zbiór testowy E1)
DISPLAY = {"paper.py", "ms/views.py"}   # warstwa wyświetlania (+ dashboard/ w etapie 4C)


def test_forecast_code_never_reads_cause_tables():
    """Zasada PIT 4N: etykiety przyczyn to wiedza po fakcie — żaden moduł poza 4N nie może ich czytać."""
    bad = []
    for p in (ROOT / "src").rglob("*.py"):
        rel = p.relative_to(ROOT / "src").as_posix()
        if rel not in ALLOWED and CAUSE_TABLES.search(p.read_text(encoding="utf-8")):
            bad.append(rel)
        if rel not in ALLOWED and re.search(r"\bfrom ms(\.| import )(causes|cause_sources)\b|import ms\.causes", p.read_text(encoding="utf-8")):
            bad.append(rel + " (import)")
    assert bad == [], f"moduły sięgają po etykiety przyczyn: {bad}"


def test_only_display_layer_imports_views():
    """ms.views czyta etykiety przyczyn do wyświetlania — kod prognoz nie może go importować (furtka obok ALLOWED)."""
    bad = []
    for p in (ROOT / "src").rglob("*.py"):
        rel = p.relative_to(ROOT / "src").as_posix()
        if rel in DISPLAY or rel.startswith("dashboard/"):
            continue
        if re.search(r"\bfrom ms(\.| import )(\w+, )*views\b|import ms\.views|from ms\.views", p.read_text(encoding="utf-8")):
            bad.append(rel)
    assert bad == [], f"ms.views importowane poza warstwą wyświetlania: {bad}"


# --------------------------------------------------------------------------- parsery

EIA = b"""<H2><A name=a1990>1990</A></H2><DL>
<DT><B>Aug</B> <DD>Iraq invades Kuwait. <I>(WSJ)</I><BR>
<DT> <DD>Conflicting reports of OPEC output.<BR>
<DT><B>Aug 2</B> <DD>raq invades Kuwait.<BR>
<DT><B>Spring</B> <DD>Seasonal note.<BR>
<DT><B>Oct. 9</B> <DD>Fear of war.</DD></DL>
<H2><A name=a1991>1991</A></H2><DL><DT><B>Jan 17</B> <DD>Operation Desert Storm begins.</DD></DL>
<P>2004 Sources include: <DL><DT><B>Jan 5</B><DD>should not be read</DD></DL>"""


def test_parse_eia_chron():
    d = cs.parse_eia_chron(EIA)
    assert list(d["event_date"].astype(str)) == ["1990-08-01", "1990-08-01", "1990-08-02", "1990-10-09", "1991-01-17"]
    assert list(d["date_precision"]) == ["month", "month", "day", "day", "day"]
    assert d["text"].iloc[0] == "Iraq invades Kuwait."            # odsyłacz (WSJ) usunięty
    assert d["label"].iloc[1] == "(ciąg dalszy)"                  # <DT> bez daty dziedziczy poprzednią
    assert "should not be read" not in " ".join(d["text"])         # sekcja po ostatnim </DL> roku pominięta
    assert d["event_id"].is_unique


def test_parse_wiki_chron():
    w = ("===1973===\n*'''October 17''': [[OPEC|Arab members]] start an [[1973 oil crisis|embargo]]<ref>x</ref>.\n"
         "*'''November''': Prices quadruple{{cn}}.\n==1974==\n===1974===\n* not bold\n").encode()
    d = cs.parse_wiki_chron(w)
    assert list(d["event_date"].astype(str)) == ["1973-10-17", "1973-11-01"]
    assert d["text"].iloc[0] == "Arab members start an embargo."
    assert list(d["date_precision"]) == ["day", "month"]


def test_parse_surprises_aggregate_same_day():
    x = io.BytesIO()
    pd.DataFrame({"Date": pd.to_datetime(["1991-02-01", "1991-02-01", "2023-12-13"]), "Time": ["9:15am", "11:30am", "2:00pm"],
                  "Unscheduled": [1, 1, 0], "MPS": [-0.04, -0.03, -0.06]}).to_excel(x, sheet_name="FOMC (update 2023)", index=False)
    s = cs.parse_sffed(x.getvalue())
    assert len(s) == 2 and s["event_id"].is_unique
    assert abs(s["value"].iloc[0] + 0.07) < 1e-12 and "suma 2" in s["text"].iloc[0]
    j = pd.DataFrame({"start": ["2005-05-03 14:15", "2005-05-03 16:00"], "MP_median": [0.004, -0.038], "CBI_median": [0.02, 0.017]})
    k = cs.parse_jk(j.to_csv(index=False).encode())
    assert len(k) == 2 and set(k["category"]) == {"fomc_mp", "fomc_cbi"}
    assert abs(k.loc[k.category == "fomc_mp", "value"].iloc[0] + 0.034) < 1e-12


def test_parse_defillama():
    js = json.dumps([{"date": 1647475200, "name": "A", "classification": "Gov", "technique": "Flash", "amount": 2e8,
                      "chain": ["Ethereum"], "bridgeHack": False, "source": "", "defillamaId": "1"},
                     {"date": 1647475200, "name": "A", "classification": "Gov", "technique": "Flash", "amount": 2e8,
                      "chain": ["Ethereum"], "bridgeHack": False, "source": "", "defillamaId": "1"}]).encode()
    d = cs.parse_defillama(js)
    assert len(d) == 1 and d["event_date"].iloc[0] == date(2022, 3, 17) and d["value"].iloc[0] == 2e8


def test_manual_and_taxonomy_files():
    m = cs.load_manual()
    assert m["event_id"].is_unique and m["verified"].notna().all()
    assert set(m["category"]) <= set(causes.taxonomy_codes())
    halv = m[m.category == "crypto_supply"]
    assert list(halv["event_date"].astype(str)) == ["2012-11-28", "2016-07-09", "2020-05-11", "2024-04-20"]
    codes = causes.taxonomy_codes()
    assert len(codes) == 20 and len(set(codes)) == 20 and "crypto_supply" in codes  # D4


# --------------------------------------------------------------------------- duże ruchy i dopasowanie


def _pf(n=400, seed=0):
    rng = np.random.default_rng(seed)
    idx = pd.bdate_range("2020-01-01", periods=n)
    r = rng.normal(0, 0.01, n)
    r[300] = 0.08
    known = idx + pd.Timedelta(hours=21)
    return pd.DataFrame({"value": 100 * np.exp(np.cumsum(r)), "logret": r, "known_at": known}, index=idx)


def test_big_moves_sigma_uses_only_past():
    pf = _pf()
    mv = causes.big_moves(pf)
    d300 = pf.index[300].date()
    assert d300 in set(mv["day"])
    s0 = mv.set_index("day").loc[d300, "sigma"]
    pf2 = pf.copy()
    pf2.iloc[300, pf2.columns.get_loc("logret")] = 0.2           # zmiana dnia D nie zmienia σ dla D
    assert causes.big_moves(pf2).set_index("day").loc[d300, "sigma"] == s0
    assert mv["day"].min() >= pf.index[252].date()                 # σ wymaga pełnych 252 sesji


def test_match_day0_day1_and_rank():
    pf = _pf()
    mv = causes.big_moves(pf)
    d300 = pf.index[300]
    ev = pd.DataFrame([
        causes._ev(d300 + pd.Timedelta(hours=14), "fomc", "monetary", ("XAU",), "fomc:x", "e", "pd"),     # przed zamknięciem -> dzień 0
        causes._ev(d300 + pd.Timedelta(hours=21), "opec", "oil_supply", ("XAU",), "opec:x", "e", "pd"),   # = known_at -> NIE ten dzień
        causes._ev(pf.index[299] + pd.Timedelta(hours=22), "gpr", "sovmil", ("XAU",), "gpr:x", "e", "pd"),  # po zamknięciu D−1 -> dzień 0
        causes._ev(d300 + pd.Timedelta(hours=14), "hack", "crypto_hack", ("BTC",), "hack:x", "e", "pd"),  # inne aktywo
    ])
    m = causes.match(pf, mv, ev, "XAU")
    got = {(r.source_ref, r.offset, r.rank, r.confidence) for r in m.itertuples() if r.day == d300.date()}
    assert got == {("fomc:x", 0, 1, 1.0), ("gpr:x", 0, 2, 1.0)}
    pf2 = pf.copy()
    pf2.iloc[301, pf2.columns.get_loc("logret")] = 0.09
    m2 = causes.match(pf2, causes.big_moves(pf2), ev, "XAU")
    got2 = {(r.source_ref, r.offset, r.confidence) for r in m2.itertuples() if r.day == pf.index[301].date()}
    assert got2 == {("opec:x", 0, 1.0), ("fomc:x", 1, 0.6), ("gpr:x", 1, 0.6)}   # OPEC: dzień 0 = D+1


# --------------------------------------------------------------------------- arkusz E2


def test_e2_sheet_stratified_prefix_and_blank():
    con = duckdb.connect()
    con.execute(db.SCHEMA)
    rows = []
    for y in range(1965, 2008):
        for k in range(12):
            rows.append(("eia_chron", f"e-{y}-{k}", date(y, k + 1, 1), "day", None, None, None, f"t{y}{k}", None, None, "pd", None))
    con.executemany("INSERT INTO cause_events VALUES (?,?,?,?,?,?,?,?,?,?,?,?)", rows)
    s = causes.e2_sheet(con)
    yrs = s["data"].str[:4].astype(int)
    assert len(s) == 150 and yrs.between(1970, 2003).all() and s["id"].is_unique
    dec = (yrs // 10 * 10).clip(upper=2000)
    full = dec.value_counts(normalize=True)
    head = dec.iloc[:100].value_counts(normalize=True)
    assert (head - full).abs().max() < 0.05                      # pierwsze 100 też jest warstwowe
    cause_col = [c for c in s.columns if c.startswith("przyczyna [")]
    assert len(cause_col) == 1 and (s[cause_col[0]] == "").all()  # pusto — bez wyników modelu
    assert "crypto_supply" in cause_col[0]


def test_describe_day_candidate_never_shown_as_cause():
    con = duckdb.connect()
    con.execute(db.SCHEMA)
    row = ("XAU", date(2013, 4, 15), -0.09, 0.009, -10.5, "sovmil", 1, "calendar", "aigpr:x", "skok AI-GPR", "", None, "",
           1.0, False, "x", "r", pd.Timestamp("2026-10-06"), "kandydat")
    con.execute("INSERT INTO move_causes VALUES (" + ",".join("?" * len(row)) + ")", list(row))
    s = causes.describe_day(con, "XAU", date(2013, 4, 15))
    assert s.startswith("tego dnia było: skok AI-GPR") and "przyczyna:" not in s
    con.execute("INSERT INTO move_causes VALUES (" + ",".join("?" * len(row)) + ")",
                list(row[:5]) + ["metal_shortage", 1, "llm", "q:1", "nagłówki", "m", "v1", "h", 0.9, True, "x", "r",
                                 pd.Timestamp("2026-10-07"), "potwierdzona"])
    s = causes.describe_day(con, "XAU", date(2013, 4, 15))
    assert s.startswith("przyczyna: metal_shortage") and "tego dnia było: skok AI-GPR" in s
    assert causes.describe_day(con, "XAU", date(2013, 4, 16)).startswith("przyczyna nieznana")


def test_e2_xlsx_roundtrip(tmp_path):
    from openpyxl import load_workbook

    csv = tmp_path / "e2.csv"
    codes = "|".join(causes.taxonomy_codes())
    pd.DataFrame({"lp": [1, 2, 3], "data": ["1973-10-17"] * 3, "dokladnosc_daty": ["day"] * 3, "zrodlo": ["eia_chron"] * 3,
                  "id": ["a", "b", "c"], "tekst": ["embargo", "OPEC cut", "pipeline"],
                  f"przyczyna [{codes}]": "", "niejasne [tak/puste]": ""}).to_csv(csv, index=False)
    x = tmp_path / "e2.xlsx"
    assert causes.e2_xlsx(csv, x) == 3
    wb = load_workbook(x)
    ws = wb[causes.E2_SHEET]
    dvs = {str(d.sqref): d.formula1 for d in ws.data_validations.dataValidation}
    assert dvs["G2:G4"] == "=kody!$A$2:$A$21" and dvs["H2:H4"] == '"tak"'
    assert ws["F2"].alignment.wrap_text and ws.column_dimensions["F"].width >= 80
    assert [c.value for c in wb["kody"]["A"]][1:] == causes.taxonomy_codes()
    assert len(causes.load_e2_gold(x, csv)) == 0
    ws["G2"], ws["H3"] = "sovmil", "tak"
    wb.save(x)
    g = causes.load_e2_gold(x, csv)
    assert list(g["id"]) == ["a", "b"] and list(g["przyczyna"]) == ["sovmil", ""]
    ws["G4"] = "wojna"                                            # kod spoza listy
    wb.save(x)
    with pytest.raises(ValueError, match="nieznane kody"):
        causes.load_e2_gold(x, csv)
    ws["G4"], ws["F3"] = None, "OPEC cut!"                        # zmieniony tekst
    wb.save(x)
    with pytest.raises(ValueError, match="tekst wpisu"):
        causes.load_e2_gold(x, csv)
