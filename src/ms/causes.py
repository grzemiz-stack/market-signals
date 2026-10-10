"""Etap 4N: duże ruchy i etykiety przyczyn z kalendarza (bez LLM). Plan: docs/ETAP4N-PLAN.md, pkt 4.

Zasada PIT: move_causes to wiedza PO FAKCIE — tylko do wyjaśniania i badań zdarzeń, nigdy cecha prognozy
(pilnuje tests/test_causes.py w bramce).

Duży ruch: |log-zwrot| > 2,5 σ, σ = odchylenie std. z 252 POPRZEDNICH sesji (bez dnia D).
Kalendarz: zdarzenie znane w chwili T -> dzień 0 = pierwsze zamknięcie aktywa znane po T (jak ms.hyp_events.car_table).
Ruch w dniu 0 -> confidence 1,0; w dniu +1 -> 0,6. Kilka trafień -> kilka wierszy; rank 1 = najwcześniejszy dzień,
potem priorytet typu zdarzenia (kolejność PRIORITY).

Etykieta kalendarza to tylko KANDYDAT (status='kandydat', validated=false): „tego dnia było: …”, nie przyczyna.
Przyczynę potwierdza tekst — zwalidowany Qwen albo człowiek (status='potwierdzona'). Pokazywanie: describe_day().
"""

from __future__ import annotations

import hashlib
import json

import numpy as np
import pandas as pd

from ms import db
from ms.cause_sources import REF
from ms.common import latest_raw
from ms.decompose import report_surprises
from ms.hyp_events import spikes
from ms.prices import ASSETS, price_frame

Z_MIN = 2.5
SIGMA_WIN = 252
CONF = {0: 1.0, 1: 0.6}
HACK_MIN_USD = 100e6
WPSR_TOP = 0.90  # niespodzianka zapasów w górnych 10% |x| względem WCZEŚNIEJSZYCH raportów
OIL, METALS, CRYPTO = ("WTI", "BRENT"), ("XAU", "XAG"), ("BTC", "ETH")
PRIORITY = ["fomc", "opec", "wpsr", "hurricane", "manual", "hack", "gpr"]
DEFS = {"z_min": Z_MIN, "sigma_win": SIGMA_WIN, "conf": CONF, "hack_min_usd": HACK_MIN_USD, "wpsr_top": WPSR_TOP,
        "gpr": "AI-GPR ACTS: skok > 95. percentyla log(x/mediana 90 dni), declustering 20 dni; gazety z dnia d -> dzień 0 = d−1",
        "priority": PRIORITY, "calendar_status": "kandydat (validated=false)"}


def definitions_hash() -> str:
    return hashlib.sha256(json.dumps(DEFS, sort_keys=True, default=str).encode()).hexdigest()[:16]


# --------------------------------------------------------------------------- duże ruchy


def big_moves(pf: pd.DataFrame, z_min: float = Z_MIN) -> pd.DataFrame:
    """pf: wynik price_frame. Zwraca dni z |z| > z_min: day, ret, sigma, z, known_at (zamknięcia)."""
    r = pf["logret"]
    sigma = r.shift(1).rolling(SIGMA_WIN, min_periods=SIGMA_WIN).std()
    z = r / sigma
    sel = z.abs() > z_min
    return pd.DataFrame({"day": pf.index[sel].date, "ret": r[sel].to_numpy(), "sigma": sigma[sel].to_numpy(),
                         "z": z[sel].to_numpy(), "known_at": pf["known_at"][sel].to_numpy()})


# --------------------------------------------------------------------------- zdarzenia kalendarza


def _ev(T, kind, cause, assets, ref, evidence, license):
    return {"T": pd.Timestamp(T), "kind": kind, "cause": cause, "assets": assets, "ref": ref, "evidence": evidence,
            "license": license}


def calendar_events(con) -> pd.DataFrame:
    """Wszystkie zdarzenia kalendarza: T (chwila, od której znane), kind, cause, assets, ref, evidence, license."""
    ev = []
    surpr = con.execute("SELECT event_date, category, value FROM cause_events WHERE source IN ('sffed','jk')").df()
    sp = {(r.event_date, r.category): r.value for r in surpr.itertuples()}
    for r in con.execute("SELECT * FROM events_calendar ORDER BY known_at").df().itertuples():
        d = pd.Timestamp(r.event_date).date()
        if r.event_type == "fomc":
            parts = [f"{lab} {sp[(d, c)]:+.3f}" for c, lab in (("fomc_mps", "MPS"), ("fomc_mp", "JK MP"), ("fomc_cbi", "JK CBI"))
                     if (d, c) in sp and pd.notna(sp[(d, c)])]
            ev.append(_ev(r.known_at, "fomc", "monetary", tuple(ASSETS), f"fomc:{d}",
                          "komunikat FOMC" + (f" ({', '.join(parts)})" if parts else " (brak miary niespodzianki)"), r.license))
        elif r.event_type == "opec":
            ev.append(_ev(r.known_at, "opec", "oil_supply", OIL, f"opec:{d}", f"komunikat OPEC (niespodzianka Känzig {r.value:+.2f})", r.license))
        elif r.event_type == "hurricane_gom":
            ev.append(_ev(r.known_at, "hurricane", "oil_supply", OIL, f"hurricane:{r.event_id}", r.note, r.license))
    fr, s = report_surprises(con)
    x = s.to_numpy()
    for i in range(len(x)):
        past = np.abs(x[:i][~np.isnan(x[:i])])
        if np.isnan(x[i]) or len(past) < 52:
            continue
        if abs(x[i]) > np.quantile(past, WPSR_TOP):
            ev.append(_ev(fr["known_at"].iloc[i], "wpsr", "oil_supply", OIL, f"wpsr:{fr['observed_at'].iloc[i].date()}",
                          f"WPSR: niespodzianka zapasów ropy {x[i]:+,.0f} tys. bbl (górne 10%)", "public_domain"))
    for r in con.execute("SELECT * FROM cause_events WHERE source='manual'").df().itertuples():
        ev.append(_ev(pd.Timestamp(r.event_date), "manual", r.category, tuple(r.assets.split(",")), f"manual:{r.event_id}",
                      r.text + ("" if str(r.verified).startswith("tak") else f" [data: {r.verified}]"), r.license))
    for r in con.execute("SELECT * FROM cause_events WHERE source='defillama' AND value >= ?", [HACK_MIN_USD]).df().itertuples():
        ev.append(_ev(pd.Timestamp(r.event_date), "hack", "crypto_hack", CRYPTO, f"defillama:{r.event_id}",
                      f"hack {r.value / 1e6:,.0f} mln USD: {r.text}", r.license))
    ev += gpr_events()
    return pd.DataFrame(ev)


def gpr_events() -> list[dict]:
    files = latest_raw("causes", "ai_gpr_data_daily.csv")
    if not files:
        return []
    g = pd.read_csv(files[-1], parse_dates=["Date"]).set_index("Date")
    out = []
    for col, assets, lab in (("GPR_OIL_ACTS", OIL, "Oil GPR (akty)"), ("ACTS_GPR_AI", METALS, "AI-GPR (akty)")):
        for t in spikes(g[col].astype(float)):
            day = (t - pd.Timedelta(days=1)).normalize()      # spikes() zwraca koniec dnia gazet d
            T = day - pd.Timedelta(days=1)                     # zdarzenie z dnia d−1 -> dzień 0 = zamknięcie d−1
            v = g[col].get(day, np.nan)
            out.append(_ev(T, "gpr", "sovmil", assets, f"aigpr:{col}:{day.date()}",
                           f"skok {lab} w gazetach z {day.date()} ({v:.0f})", "academic (niepotwierdzona)"))
    return out


# --------------------------------------------------------------------------- dopasowanie


def match(pf: pd.DataFrame, moves: pd.DataFrame, events: pd.DataFrame, asset: str) -> pd.DataFrame:
    """Dla zdarzeń dotyczących aktywa: dzień 0 i +1 wg known_at zamknięć; zwraca trafienia w duże ruchy."""
    p = pf[pf["known_at"].notna()].sort_values("known_at")
    known = p["known_at"].to_numpy(dtype="datetime64[ns]")
    days = p.index.to_numpy()
    mv = moves.set_index("day")
    rows = []
    for e in events[events["assets"].map(lambda a: asset in a)].itertuples():
        d0 = np.searchsorted(known, np.datetime64(e.T), side="right")
        for off in (0, 1):
            if d0 + off >= len(days):
                break
            day = pd.Timestamp(days[d0 + off]).date()
            if day in mv.index:
                m = mv.loc[day]
                rows.append({"asset": asset, "day": day, "ret": m.ret, "sigma": m.sigma, "z": m.z, "cause": e.cause,
                             "kind": e.kind, "offset": off, "source_ref": e.ref, "evidence": e.evidence,
                             "confidence": CONF[off], "license": e.license})
    out = pd.DataFrame(rows)
    if out.empty:
        return out
    out["prio"] = out["kind"].map(PRIORITY.index)
    out = out.sort_values(["day", "offset", "prio"])
    out["rank"] = np.where(out.groupby("day").cumcount() == 0, 1, 2)
    return out.drop(columns=["prio"])


def label_calendar(con, run_id: str) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Zapis etykiet kalendarza do move_causes (idempotentnie). Zwraca (etykiety, duże ruchy wszystkich aktywów)."""
    events = calendar_events(con)
    labs, mvs = [], []
    for asset in ASSETS:
        pf = price_frame(con, asset)
        mv = big_moves(pf).assign(asset=asset)
        mvs.append(mv)
        labs.append(match(pf, mv, events, asset))
    lab = pd.concat([x for x in labs if not x.empty], ignore_index=True)
    now = db.utcnow()
    lab = lab.assign(label_source="calendar", model="", prompt_version=None, prompt_hash="", validated=False,
                     run_id=run_id, known_at=now, status="kandydat")
    cols = ["asset", "day", "ret", "sigma", "z", "cause", "rank", "label_source", "source_ref", "evidence", "model",
            "prompt_version", "prompt_hash", "confidence", "validated", "license", "run_id", "known_at", "status"]
    # wiersze z pierwszego przebiegu (validated=true, przed decyzją o kandydatach) -> kandydat
    con.execute("UPDATE move_causes SET validated = false, status = 'kandydat' "
                "WHERE label_source = 'calendar' AND (validated OR status IS DISTINCT FROM 'kandydat')")
    con.register("_mc", lab[cols])
    con.execute(f"INSERT OR IGNORE INTO move_causes ({', '.join(cols)}) SELECT {', '.join(cols)} FROM _mc")  # tylko dopisywanie
    con.unregister("_mc")
    return lab, pd.concat(mvs, ignore_index=True)


def describe_day(con, asset: str, day) -> str:
    """Tekst do statusu/raportów. Potwierdzone -> „przyczyna: …”; kandydaci z kalendarza -> „tego dnia było: …”.
    Kandydat nigdy nie jest pokazywany jako przyczyna."""
    d = con.execute("""SELECT cause, evidence, status, rank, confidence FROM move_causes WHERE asset = ? AND day = ?
                       ORDER BY status = 'potwierdzona' DESC, rank, confidence DESC""", [asset, day]).df()
    if d.empty:
        return "przyczyna nieznana (brak zdarzeń w kalendarzu)"
    conf = d[d.status == "potwierdzona"]
    cand = d[d.status != "potwierdzona"].drop_duplicates("evidence")
    parts = []
    if len(conf):
        parts.append("przyczyna: " + "; ".join(f"{r.cause} ({r.evidence})" for r in conf.drop_duplicates("cause").itertuples()))
    if len(cand):
        parts.append("tego dnia było: " + "; ".join(cand["evidence"]) + " — kandydat, przyczyna niepotwierdzona tekstem")
    return " | ".join(parts)


# --------------------------------------------------------------------------- arkusz E2 do oznaczenia przez użytkownika


def taxonomy_codes() -> list[str]:
    return pd.read_csv(REF / "taxonomy.csv", comment="#")["code"].tolist()


def e2_sheet(con, n: int = 150, seed: int = 20261006) -> pd.DataFrame:
    """150 wpisów chronologii ropy 1970–2003 (EIA + Wikipedia), losowanie warstwowe po dekadach, kolejność losowa:
    każdy początkowy fragment (np. pierwsze 100) też jest próbą warstwową. Bez żadnych wyników modelu."""
    d = con.execute("""SELECT source, event_id, event_date, date_precision, text FROM cause_events
                       WHERE source IN ('eia_chron', 'wiki_chron') AND year(event_date) BETWEEN 1970 AND 2003
                       ORDER BY source, event_id""").df()
    d["decade"] = (pd.to_datetime(d["event_date"]).dt.year // 10 * 10).clip(upper=2000)
    rng = np.random.default_rng(seed)
    share = d["decade"].value_counts(normalize=True).sort_index()
    take = (share * n).round().astype(int)
    take.iloc[-1] += n - take.sum()
    parts = [g.iloc[rng.permutation(len(g))[: take[dec]]] for dec, g in d.groupby("decade")]
    s = pd.concat(parts)
    # przeplot dekad: kolejność wg losowej pozycji względnej wewnątrz dekady
    s["k"] = s.groupby("decade").cumcount() / s["decade"].map(take) + rng.uniform(0, 1e-6, len(s))
    s = s.sort_values("k").reset_index(drop=True)
    codes = "|".join(taxonomy_codes())
    return pd.DataFrame({"lp": range(1, len(s) + 1), "data": s["event_date"].astype(str),
                         "dokladnosc_daty": s["date_precision"], "zrodlo": s["source"], "id": s["event_id"],
                         "tekst": s["text"], f"przyczyna [{codes}]": "", "niejasne [tak/puste]": ""})


# --------------------------------------------------------------------------- arkusz E2 w xlsx (oznaczanie) i wczytanie

E2_SHEET = "E2"
E2_COLS = ["lp", "data", "dokladnosc_daty", "zrodlo", "id", "tekst", "przyczyna", "niejasne"]


def e2_xlsx(csv_path, xlsx_path) -> int:
    """CSV (źródło) -> xlsx do oznaczania: lista rozwijana kodów w „przyczyna”, „niejasne” = tak/puste,
    szeroka kolumna tekstu z zawijaniem. Zakładka „kody” z opisami."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.worksheet.datavalidation import DataValidation

    src = pd.read_csv(csv_path, dtype=str, keep_default_na=False)
    tax = pd.read_csv(REF / "taxonomy.csv", comment="#", dtype=str, keep_default_na=False)
    wb = Workbook()
    ws = wb.active
    ws.title = E2_SHEET
    ws.append(E2_COLS)
    for r in src.itertuples(index=False):
        ws.append([int(r[0]), r[1], r[2], r[3], r[4], r[5], None, None])
    n = len(src) + 1
    for c in ws[1]:
        c.font = Font(bold=True)
    for col, w in zip("ABCDEFGH", (5, 11, 9, 11, 18, 100, 24, 10)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2, max_row=n):
        for c in row:
            c.alignment = Alignment(wrap_text=c.column_letter == "F", vertical="top")
    ws.freeze_panes = "A2"
    k = wb.create_sheet("kody")
    k.append(["kod", "opis"])
    for r in tax.itertuples(index=False):
        k.append([r.code, r.opis])
    k.column_dimensions["A"].width = 24
    k.column_dimensions["B"].width = 110
    dv = DataValidation(type="list", formula1=f"=kody!$A$2:$A${len(tax) + 1}", allow_blank=True,
                        showErrorMessage=True, errorTitle="Nieznany kod", error="Wybierz kod z listy (zakładka „kody”).")
    dv2 = DataValidation(type="list", formula1='"tak"', allow_blank=True, showErrorMessage=True,
                         errorTitle="niejasne", error="Wpisz „tak” albo zostaw puste.")
    ws.add_data_validation(dv)
    ws.add_data_validation(dv2)
    dv.add(f"G2:G{n}")
    dv2.add(f"H2:H{n}")
    i = wb.create_sheet("instrukcja")
    for line in ("Oznacz przyczynę ruchu ceny ropy, którą opisuje wpis (kolumna „przyczyna”, lista rozwijana).",
                 "Opisy kodów: zakładka „kody”. Wpis niezwiązany z ceną ropy albo bez wyjaśnienia -> kod „unknown”.",
                 "Gdy nie da się zdecydować -> „niejasne” = tak (przyczyna może zostać pusta).",
                 "Wiersze są w losowej kolejności; wystarczy oznaczyć pierwsze 100 (każdy początek jest próbą warstwową).",
                 "Nie zmieniaj kolumn lp, id ani tekstu — wczytanie porównuje je z CSV i odrzuci arkusz przy różnicy."):
        i.append([line])
    i.column_dimensions["A"].width = 120
    wb.save(xlsx_path)
    return len(src)


def load_e2_gold(xlsx_path, csv_path) -> pd.DataFrame:
    """Wczytuje oznaczony xlsx. Sprawdza zgodność z CSV (id, tekst, kolejność) i kody. Zwraca wiersze oznaczone
    (przyczyna albo niejasne=tak) z kolumnami: lp, id, tekst, przyczyna, niejasne."""
    x = pd.read_excel(xlsx_path, sheet_name=E2_SHEET, dtype=str, keep_default_na=False)
    src = pd.read_csv(csv_path, dtype=str, keep_default_na=False)
    errs = []
    if list(x.columns) != E2_COLS:
        errs.append(f"kolumny {list(x.columns)} != {E2_COLS}")
    elif len(x) != len(src) or (x["id"].to_numpy() != src["id"].to_numpy()).any():
        errs.append("id / kolejność wierszy różni się od CSV")
    elif (x["tekst"].to_numpy() != src["tekst"].to_numpy()).any():
        errs.append("tekst wpisu różni się od CSV: " + ", ".join(x.loc[x["tekst"] != src["tekst"], "lp"].astype(str)))
    if not errs:
        x["przyczyna"] = x["przyczyna"].str.strip()
        x["niejasne"] = x["niejasne"].str.strip().str.lower()
        codes = set(taxonomy_codes())
        bad = x[(x["przyczyna"] != "") & ~x["przyczyna"].isin(codes)]
        if len(bad):
            errs.append("nieznane kody: " + ", ".join(f"lp {r.lp}: {r.przyczyna}" for r in bad.itertuples()))
        bad = x[~x["niejasne"].isin(["", "tak"])]
        if len(bad):
            errs.append("„niejasne” poza tak/puste: lp " + ", ".join(bad["lp"]))
    if errs:
        raise ValueError("; ".join(errs))
    out = x[(x["przyczyna"] != "") | (x["niejasne"] == "tak")]
    return out[["lp", "id", "tekst", "przyczyna", "niejasne"]].assign(lp=lambda d: d["lp"].astype(int)).reset_index(drop=True)
